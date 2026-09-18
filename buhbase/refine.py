# -*- coding: utf-8 -*-
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

src = openpyxl.load_workbook("Бухгалтеры_Барнаул_пилот.xlsx").active
rows = list(src.iter_rows(min_row=2, values_only=True))

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "База A-B-C"
hdr = ["Сегмент","Качество","Осн. контакт","Email","Организация / ФИО","Город",
       "Профиль","Источник (2ГИС)","Сайт","Пометка","Дата сбора"]
ws.append(hdr)

core = 0
for seg, email, org, city, prof, s_url, site, note, date in rows:
    note = note or ""
    disputed = "(?)" in note
    secondary = ("второй адрес" in note)
    quality = "спорная" if disputed else "ядро"
    main = "нет" if secondary else "да"
    if not disputed and not secondary:
        core += 1
    ws.append([seg, quality, main, email, org, city, prof, s_url, site, note, date])

hf = Font(bold=True, color="FFFFFF")
hfill = PatternFill("solid", fgColor="2F5597")
for c in ws[1]:
    c.font = hf; c.fill = hfill
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
ws.freeze_panes = "A2"
ws.auto_filter.ref = ws.dimensions

seg_fill = {"A": PatternFill("solid", fgColor="FFF2CC"),
            "B": PatternFill("solid", fgColor="DEEBF7"),
            "C": PatternFill("solid", fgColor="E2EFDA")}
grey = Font(color="808080")
warn = PatternFill("solid", fgColor="FCE4D6")
for row in ws.iter_rows(min_row=2):
    if row[0].value in seg_fill:
        row[0].fill = seg_fill[row[0].value]
    row[0].alignment = Alignment(horizontal="center")
    row[1].alignment = Alignment(horizontal="center")
    row[2].alignment = Alignment(horizontal="center")
    if row[1].value == "спорная":
        row[1].fill = warn
    if row[2].value == "нет":
        for c in row:
            c.font = grey

for i, w in enumerate([9, 11, 9, 30, 46, 15, 30, 46, 24, 46, 12], 1):
    ws.column_dimensions[get_column_letter(i)].width = w

# лист со сводкой
s = wb.create_sheet("Сводка")
from collections import Counter
seg_c = Counter(r[0] for r in rows)
qual_c = Counter("спорная" if "(?)" in (r[7] or "") else "ядро" for r in rows)
sec_c = sum(1 for r in rows if "второй адрес" in (r[7] or ""))
orgs = len({r[2] for r in rows})

s.append(["Показатель", "Значение"])
for k, v in [("Всего уникальных email", len(rows)),
             ("Уникальных организаций", orgs),
             ("Сегмент A — частные бухгалтеры", seg_c["A"]),
             ("Сегмент B — бухгалтерии компаний", seg_c["B"]),
             ("Сегмент C — бухгалтерские компании", seg_c["C"]),
             ("Качество: ядро", qual_c["ядро"]),
             ("Качество: спорная", qual_c["спорная"]),
             ("Вторые адреса тех же компаний", sec_c),
             ("Твёрдое ядро (ядро + осн. контакт)", core),
             ("География", "Барнаул, Бийск, Горно-Алтайск"),
             ("Источники", "2ГИС + сайты организаций"),
             ("Дата сбора", "2026-09-18")]:
    s.append([k, v])
for c in s[1]:
    c.font = hf; c.fill = hfill
s.column_dimensions["A"].width = 40
s.column_dimensions["B"].width = 32
s.freeze_panes = "A2"

wb.save("Бухгалтеры_Барнаул_пилот.xlsx")
print("строк:", ws.max_row - 1, "| ядро+осн:", core, "| организаций:", orgs)
