# -*- coding: utf-8 -*-
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from collections import Counter

D = "2026-09-18"
B = "https://2gis.ru/barnaul/firm/"
BI = "https://2gis.ru/biysk/firm/"
GA = "https://2gis.ru/gornoaltaysk/firm/"
NA = "https://2gis.ru/novoaltajsk/firm/"
RU = "https://2gis.ru/rubtsovsk/firm/"

# исходные 52 записи пилота
old = openpyxl.load_workbook("Бухгалтеры_Барнаул_пилот.xlsx")["База A-B-C"]
base = []
for r in old.iter_rows(min_row=2, values_only=True):
    seg, q, m, email, org, city, prof, src, site, note, date = r
    base.append((seg, email, org, city, prof, src, site, note or ""))

# новые записи второй волны
new = [
# --- Сегмент B: бухгалтерии внутри компаний ---
("B","glav.buh@z-km.com","Котломаш, производство котельного оборудования","Барнаул","Производство",B+"—","z-km.com","маркер glav.buh в адресе"),
("B","buh@cp22.ru","Строительная перспектива","Барнаул","Строительство",B+"—","cp22.ru","маркер buh@ в адресе"),
("B","glbuh@oe22.ru","Ока-Электрон","Барнаул","Производство / электроника",B+"—","oe22.ru","маркер glbuh в адресе"),
("B","finance@altai-tent.ru","Алтай-Тент","Барнаул","Производство тентов",B+"—","altai-tent.ru","маркер finance@ в адресе"),
("B","mku@cuaio.ru","Центр учёта, анализа и отчётности Майминского района (МКУ)","Майма, Респ. Алтай","Централизованная бухгалтерия",GA+"70000001104984629","cuaio.ru","муниципальное учреждение, профиль — бухучёт"),
("B","ekon999@bk.ru","Автоимперия","Барнаул","Автоперевозки",B+"—","avtoimperiya.ru","(?) адрес похож на экономический отдел, роль не подписана"),
# --- Сегмент A: частные бухгалтеры ---
("A","nata.zvereva77@bk.ru","Зверева Н. В., бухгалтерская компания","Бийск","Бухгалтерские услуги",BI+"70000001114864670","zvereva-buh.online","ФИО в названии, личный адрес"),
("A","rakshina7@bk.ru","Бухгалтерия Малого Бизнеса (Ракшина)","Бийск","Бухгалтерские услуги",BI+"2815279048498835","","(?) личный адрес, микропрактика"),
("A","christinahomyackowa@yandex.ru","Альфа, центр финансовых услуг (Хомякова)","Бийск","Бухгалтерские услуги",BI+"70000001044727624","alfabuhg.ru","(?) личный адрес специалиста на сайте"),
("A","asanova-22@mail.ru","Аудит-экономикс (Асанова)","Барнаул","Аудит / бухучёт",B+"563478234415097","audit-economics.ru","(?) личный адрес специалиста на сайте"),
# --- Сегмент C: Бийск ---
("C","mail@marex-audit.ru","Марэкс Аудит","Бийск","Аудит / бухучёт",BI+"70000001065386440","mareks-audit.ru",""),
("C","nalog-cons@bk.ru","Налоговая консультация","Бийск","Бухгалтерские услуги",BI+"70000001097383327","",""),
("C","buh555354@mail.ru","Бухсервис, бухгалтерская компания","Бийск","Бухгалтерские услуги",BI+"70000001021751684","","маркер buh в адресе"),
("C","preventiva@list.ru","Превентива Плюс, консалтинговое объединение","Бийск","Бухгалтерские услуги",BI+"2815278048091386","biysk.preventiva.ru",""),
("C","ktk_biysk@mail.ru","Бухгалтерский центр","Бийск","Бухгалтерские услуги",BI+"2815278048091378","buhbiysk.ru",""),
("C","tend.nalog@mail.ru","Декларант, многопрофильная компания","Бийск","Бухгалтерские услуги",BI+"70000001021088443","deklarant-sibir.ru",""),
("C","1cbo@me-72.ru","БухПомощь, бухгалтерская компания","Бийск","1С:БухОбслуживание",BI+"70000001111480962","1cbo.helpbuh.online",""),
("C","marion.w@mail.ru","Марион, бухгалтерская фирма","Бийск","Бухгалтерские услуги",BI+"2815278048266652","marion22.ru",""),
("C","info@adaru22.ru","Адару, консалтинговая группа","Бийск","Бухгалтерские услуги",BI+"2815278048113296","adaru22.ru",""),
("C","alfa.buhgalteriya@mail.ru","Альфа, центр финансовых услуг","Бийск","Бухгалтерские услуги",BI+"70000001044727624","alfabuhg.ru","маркер buhgalteriya в адресе"),
("C","Experta22@yandex.ru","Эксперта","Бийск","Бухгалтерские услуги",BI+"70000001042885519","22-experta.ru",""),
("C","cnp.bsk@bk.ru","Центр налогового планирования","Бийск","Бухгалтерские услуги",BI+"70000001043532101","cnp-nalog.ru",""),
("C","bravoltd.22@gmail.com","Браво, бухгалтерская компания","Бийск","Бухгалтерские услуги",BI+"70000001032344568","bravoltd.8b.io",""),
# --- Сегмент C: Горно-Алтайск ---
("C","info@tekabuh.ru","Тека бухгалтерия, бухгалтерская компания","Горно-Алтайск","Бухгалтерские услуги",GA+"3800440466595925","tekabuh.ru",""),
("C","glavbuh04@bk.ru","Тека бухгалтерия, бухгалтерская компания","Горно-Алтайск","Бухгалтерские услуги",GA+"3800440466595925","tekabuh.ru","второй адрес той же компании, маркер glavbuh"),
("C","vash_buhgalter_04@mail.ru","БизнесЭксперТ, центр сопровождения бизнеса","Горно-Алтайск","Бухгалтерские услуги",GA+"3800440466582934","center-bet.ru","маркер buhgalter в адресе"),
("C","buscenter05@mail.ru","БизнесЭксперТ, центр сопровождения бизнеса","Горно-Алтайск","Бухгалтерские услуги",GA+"3800440466582934","center-bet.ru","второй адрес той же компании"),
("C","finuspeh04@mail.ru","ФинУспех, бухгалтерская компания","Горно-Алтайск","Бухгалтерские услуги",GA+"70000001030140737","finuspeh04.ru",""),
("C","3818506@mail.ru","Гарантия, юридический центр Светляковой","Горно-Алтайск","Юруслуги / бухучёт",GA+"70000001051290223","","(?) профиль ближе к банкротству"),
# --- Сегмент C: Новоалтайск ---
("C","nalog.sovetnik@mail.ru","Налоговый советник, бухгалтерско-правовая компания","Новоалтайск","Бухгалтерские услуги",NA+"70000001019454980","nalogsovetnik.ru",""),
# --- Сегмент C: Барнаул, вторая волна ---
("C","kostar.buh@yandex.ru","Костар, бухгалтерская аутсорсинговая компания","Барнаул","Бухгалтерские услуги",B+"70000001084517439","костар.рф","маркер buh в адресе"),
("C","Zorinatt@mail.ru","Новые решения, центр сопровождения бизнеса","Барнаул","Бухгалтерские услуги",B+"70000001037273296","новые-решения.com",""),
("C","statusprofbuh@gmail.com","СтатусПроф, бухгалтерская компания","Барнаул","Бухгалтерские услуги",B+"70000001032724575","статуспроф.рф","маркер buh в адресе"),
("C","mail@uchetbezzabot.ru","Учёт без забот","Барнаул","Бухгалтерские услуги / курсы",B+"70000001082683155","учетбеззабот.рф",""),
("C","trans-audit2009@mail.ru","Транс-аудит, аудиторская фирма","Барнаул","Аудиторские услуги",B+"563478235281112","транс-аудит.рф",""),
("C","ada-audit@mail.ru","Алтайский дом аудита","Барнаул","Аудит / бухучёт",B+"563478234400898","ada-audit.ru",""),
("C","vash-yurist22@mail.ru","Ваш бухгалтер и юрист, аутсорсинговая компания","Барнаул","Бухгалтерские услуги",B+"70000001040379977","ваш-бухгалтер-и-юрист.рф",""),
("C","auditec22@mail.ru","Аудит-экономикс, аудиторская компания","Барнаул","Аудит / бухучёт",B+"563478234415097","audit-economics.ru",""),
("C","ovk_audit@mail.ru","ОВК-Аудит, аудиторская фирма","Барнаул","Аудиторские услуги",B+"563478234750377","ovk-audit.ru",""),
("C","konturs@internet.ru","Контур Сибирь / СБИС ЭПД","Барнаул","Бухгалтерские услуги / ЭДО",B+"70000001114960522","эдо.site","(?) профиль ближе к ЭДО"),
("C","690985@mail.ru","RestoFlow либо Советник (домен tilda.ws)","Барнаул","Бухгалтерские услуги",B+"—","*.tilda.ws","(?) поддомен схлопнулся, принадлежность не подтверждена"),
]

rows = base + new

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "База A-B-C"
hdr = ["Сегмент","Качество","Осн. контакт","Email","Организация / ФИО","Город",
       "Профиль","Источник (2ГИС)","Сайт","Пометка","Дата сбора"]
ws.append(hdr)

seen, dup, core = set(), 0, 0
for seg, email, org, city, prof, src, site, note in rows:
    k = email.lower().strip()
    if k in seen:
        dup += 1
        continue
    seen.add(k)
    disputed = "(?)" in note
    secondary = "второй адрес" in note
    quality = "спорная" if disputed else "ядро"
    main = "нет" if secondary else "да"
    if not disputed and not secondary:
        core += 1
    ws.append([seg, quality, main, email, org, city, prof, src, site, note, D])

hf = Font(bold=True, color="FFFFFF"); hfill = PatternFill("solid", fgColor="2F5597")
for c in ws[1]:
    c.font = hf; c.fill = hfill
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions

seg_fill = {"A": PatternFill("solid", fgColor="FFF2CC"),
            "B": PatternFill("solid", fgColor="DEEBF7"),
            "C": PatternFill("solid", fgColor="E2EFDA")}
grey = Font(color="808080"); warn = PatternFill("solid", fgColor="FCE4D6")
for row in ws.iter_rows(min_row=2):
    if row[0].value in seg_fill: row[0].fill = seg_fill[row[0].value]
    for i in (0,1,2): row[i].alignment = Alignment(horizontal="center")
    if row[1].value == "спорная": row[1].fill = warn
    if row[2].value == "нет":
        for c in row: c.font = grey

for i, w in enumerate([9,11,9,32,50,17,32,46,26,52,12], 1):
    ws.column_dimensions[get_column_letter(i)].width = w

data = [r for r in ws.iter_rows(min_row=2, values_only=True)]
segc = Counter(r[0] for r in data)
cityc = Counter(r[5] for r in data)
orgs = len({r[4] for r in data})

s = wb.create_sheet("Сводка")
s.append(["Показатель","Значение"])
for k,v in [("Всего уникальных email", len(data)),
            ("Уникальных организаций", orgs),
            ("Сегмент A — частные бухгалтеры", segc["A"]),
            ("Сегмент B — бухгалтерии компаний", segc["B"]),
            ("Сегмент C — бухгалтерские компании", segc["C"]),
            ("Качество: ядро", sum(1 for r in data if r[1]=="ядро")),
            ("Качество: спорная", sum(1 for r in data if r[1]=="спорная")),
            ("Твёрдое ядро (ядро + осн. контакт)", core),
            ("Дубликатов отброшено", dup),
            ("Дата сбора", D)]:
    s.append([k,v])
s.append(["",""]); s.append(["Город","Записей"])
for c,n in cityc.most_common():
    s.append([c or "не указан", n])
for c in s[1]: c.font = hf; c.fill = hfill
for c in s[12]: c.font = hf; c.fill = hfill
s.column_dimensions["A"].width = 42; s.column_dimensions["B"].width = 30
s.freeze_panes = "A2"

wb.save("Бухгалтеры_Алтай_база.xlsx")
print("всего:", len(data), "| A:", segc["A"], "B:", segc["B"], "C:", segc["C"])
print("ядро+осн:", core, "| организаций:", orgs, "| дубли:", dup)
print("города:", dict(cityc))
